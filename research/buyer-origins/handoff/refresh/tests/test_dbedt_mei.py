"""Tests for the DBEDT monthly housing refresh.

Run from the refresh folder:  python3 -m unittest discover -s tests -v
Fixtures are the real DBEDT workbooks of the July 2026 edition (posted 2026-08-28) and the
August 2026 edition (released 2026-10-01), plus the publisher page as fetched on 2026-10-09.
Failure and revision cases edit copies of those files in a temporary folder.
"""
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
sys.path.insert(0, PKG)

import dbedt_mei as d  # noqa: E402
import xlsx_lite  # noqa: E402

FIX = os.path.join(HERE, 'fixtures')
AUG, JUL = os.path.join(FIX, 'aug2026'), os.path.join(FIX, 'jul2026')
BASELINE = d.load_baseline(os.path.join(PKG, 'baseline', 'verified_2026-08.json'))
URL = 'https://files.hawaii.gov/dbedt/economic/data_reports/mei/'
REAL_REVISIONS_JUL_TO_AUG = {
    ('maui', '2026-06', 'condo_median', 630000, 625000),
    ('kauai', '2025-08', 'condo_sales', 25, 26),
    ('kauai', '2025-08', 'condo_median', 675000, 680000),
    ('hawaii', '2025-08', 'sf_sales', 165, 166),
    ('hawaii', '2025-08', 'sf_median', 527000, 523500),
}
COL = {'sf_sales': 'AO', 'sf_median': 'AP', 'condo_sales': 'AR', 'condo_median': 'AS'}


def read(path):
    with open(path, 'rb') as f:
        return f.read()


def row_of(data, month):
    """1-based worksheet row holding a month in a county workbook."""
    wb = xlsx_lite.Workbook(data)
    for i, r in enumerate(wb.rows(0)):
        if r and isinstance(r[0], (int, float)) and 30000 <= r[0] <= 80000:
            dt = xlsx_lite.excel_serial_to_date(r[0])
            if '%04d-%02d' % (dt.year, dt.month) == month:
                return i + 1
    raise KeyError(month)


def patch_xlsx(data, edits, sheet_path='xl/worksheets/sheet1.xml'):
    """Copy of an .xlsx with some cells replaced (None blanks a cell; str becomes an inline string)."""
    zin, buf = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            b = zin.read(item.filename)
            if item.filename == sheet_path:
                x = b.decode('utf-8')
                for ref, val in edits.items():
                    m = re.search(r'<c r="%s"(?P<a>[^>]*?)(?:/>|>(?P<b>.*?)</c>)' % ref, x, flags=re.S)
                    assert m, 'cell %s not found' % ref
                    attrs = re.sub(r'\s+t="[^"]*"', '', m.group('a')).rstrip('/')
                    if val is None:
                        new = '<c r="%s"%s/>' % (ref, attrs)
                    elif isinstance(val, str):
                        new = '<c r="%s"%s t="inlineStr"><is><t>%s</t></is></c>' % (ref, attrs, escape(val))
                    else:
                        new = '<c r="%s"%s><v>%s</v></c>' % (ref, attrs, val)
                    x = x[:m.start()] + new + x[m.end():]
                b = x.encode('utf-8')
            zout.writestr(item, b)
    return buf.getvalue()


def july_page():
    page = read(os.path.join(AUG, 'mei_index.html')).decode('utf-8')
    page = page.replace('through August 2026 was released on October 1st', 'through July 2026 was released on August 28th')
    return page.replace('2026-08-', '2026-07-').replace('mei-08p-2026', 'mei-07p-2026').replace('August 2026', 'July 2026')


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='dbedt_test_')
        self.state = os.path.join(self.tmp, 'state')

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def edition(self, which, page=None, edit=None):
        """A folder holding one edition's page and workbooks; `edit` maps file name -> bytes transformer or None (delete)."""
        src = os.path.join(self.tmp, '%s_%d' % (which, len(os.listdir(self.tmp))))
        shutil.copytree(AUG if which == 'aug' else JUL, src)
        if which == 'jul':
            with open(os.path.join(src, 'mei_index.html'), 'w', encoding='utf-8') as f:
                f.write(july_page())
        if page is not None:
            html = page(read(os.path.join(src, 'mei_index.html')).decode('utf-8'))
            with open(os.path.join(src, 'mei_index.html'), 'w', encoding='utf-8') as f:
                f.write(html)
        for name, fn in (edit or {}).items():
            p = os.path.join(src, name)
            if fn is None:
                os.remove(p)
            else:
                new = fn(read(p))
                with open(p, 'wb') as f:
                    f.write(new)
        return src

    def run_edition(self, folder, now, **kw):
        return d.run(d.DirectoryFetcher(folder, now=now), self.state, baseline=BASELINE, now=now, **kw)

    def last_good(self):
        return d.read_json(os.path.join(self.state, 'last_good', d.DATASET_FILE))

    def last_good_bytes(self):
        return read(os.path.join(self.state, 'last_good', d.DATASET_FILE))


class ReaderAndDiscovery(Base):
    @unittest.skipUnless(__import__('importlib').util.find_spec('openpyxl'), 'openpyxl not installed')
    def test_reader_matches_openpyxl_cell_for_cell(self):
        import datetime
        import openpyxl
        for folder in (AUG, JUL):
            for name in sorted(os.listdir(folder)):
                if not name.endswith('.xlsx'):
                    continue
                p = os.path.join(folder, name)
                wb, ow = xlsx_lite.Workbook(read(p)), openpyxl.load_workbook(p, data_only=True)
                for si in range(len(wb.sheets)):
                    a, b = wb.rows(si), [list(r) for r in ow.worksheets[si].iter_rows(values_only=True)]
                    for i in range(max(len(a), len(b))):
                        ra, rb = (a[i] if i < len(a) else []), (b[i] if i < len(b) else [])
                        for j in range(max(len(ra), len(rb))):
                            x, y = (ra[j] if j < len(ra) else None), (rb[j] if j < len(rb) else None)
                            if isinstance(y, datetime.datetime):
                                y = (y.date() - datetime.date(1899, 12, 30)).days
                            self.assertEqual(x, None if y == '' else y, '%s %s r%d c%d' % (name, si, i + 1, j))

    def test_discovers_latest_complete_set_from_publisher_page(self):
        disc = d.discover(read(os.path.join(AUG, 'mei_index.html')).decode('utf-8'))
        self.assertEqual(disc['latest_complete_period'], '2026-08')
        for c in d.COUNTIES:
            self.assertEqual(disc['periods']['2026-08'][c]['url'], URL + '2026-08-%s.xlsx' % c)
        self.assertEqual(disc['summaries']['2026-08']['url'], URL + 'mei-08p-2026.xlsx')
        self.assertEqual(disc['summaries']['2026-08']['file_suffix'], 'p')
        self.assertEqual(len(disc['page_statements']), 2)
        self.assertIn('through September 2026 will be released on/about October 29th', disc['page_statements'][1])

    def test_discovery_resolves_relative_links_and_ignores_other_hosts(self):
        page = ''.join('<a href="/dbedt/economic/data_reports/mei/2026-09-%s.xlsx">%s County</a>' % (c, c) for c in d.COUNTIES)
        page = page.replace('href="/dbedt', 'href="https://files.hawaii.gov/dbedt', 3)  # three absolute, one relative
        page += '<a href="https://evil.example/mei/2026-10-honolulu.xlsx">Honolulu County</a>'
        disc = d.discover(page, page_url='https://files.hawaii.gov/dbedt/economic/')
        self.assertEqual(disc['latest_complete_period'], '2026-09')
        self.assertNotIn('2026-10', disc['periods'])
        self.assertTrue(any('non-DBEDT host' in p for p in disc['problems']))


class AugustEdition(Base):
    @classmethod
    def setUpClass(cls):
        cls.parsed = {c: d.parse_county_workbook(read(os.path.join(AUG, '2026-08-%s.xlsx' % c)), c) for c in d.COUNTIES}

    def test_every_baseline_value_matches_the_verified_report_series(self):
        n = 0
        for c in d.COUNTIES:
            self.assertEqual(self.parsed[c]['errors'], [])
            for month, vals in BASELINE['counties'][c].items():
                for f in d.FIELDS:
                    self.assertEqual(self.parsed[c]['months'].get(month, {}).get(f), vals[f], (c, month, f))
                    n += 1
        self.assertEqual(n, 3968)

    def test_missing_values_stay_missing(self):
        k = self.parsed['kauai']['months']
        self.assertIsNone(k['2006-03']['sf_sales'])           # no monthly counts Mar 2006 - Dec 2012
        self.assertIsNone(k['2006-03']['sf_median'])          # no monthly medians Mar 2006 - Dec 2007
        self.assertEqual(self.parsed['kauai']['first_month'], '2005-01')
        ytd = d.compute_ytd(k, '2026-08')
        self.assertFalse(ytd['2006']['sf_sales']['complete'])
        self.assertIsNone(ytd['2006']['sf_sales']['value'])
        self.assertTrue(ytd['2013']['sf_sales']['complete'])
        for c in d.COUNTIES:
            for v in self.parsed[c]['months'].values():
                for f in d.MEDIAN_OF:
                    self.assertNotEqual(v[f], 0)

    def test_real_odd_value_is_kept_and_small_sample_flagged(self):
        v = self.parsed['kauai']['months']['2026-08']
        self.assertEqual(v['condo_median'], 999999)   # a real median (Title Guaranty's report shows the same)
        self.assertEqual(v['condo_sales'], 15)
        self.assertEqual(d.small_sample_flags(v), ['condo_median'])

    def test_full_run_promotes_with_every_check_passing(self):
        rec = self.run_edition(AUG, '2026-10-09T02:00:00Z')
        self.assertEqual((rec['decision'], rec['exit_code'], rec['status']), ('promoted', 0, 'pass'))
        self.assertEqual([c for c in rec['checks'] if c['status'] != 'pass'], [])
        ds = self.last_good()
        self.assertEqual(ds['release']['observation_month'], '2026-08')
        self.assertEqual(ds['release']['publication_date'], '2026-10-01')
        self.assertEqual({s['role'] for s in ds['sources']}, {'publisher_page', 'county_workbook', 'summary_workbook'})
        for s in ds['sources']:
            self.assertRegex(s['sha256'], r'^[0-9a-f]{64}$')

    def test_monthly_ytd_and_statewide_are_kept_apart(self):
        self.run_edition(AUG, '2026-10-09T02:00:00Z')
        ds = self.last_good()
        h = ds['counties']['honolulu']['months']['2026-08']
        self.assertEqual((h['sf_sales'], h['sf_median'], h['condo_sales'], h['condo_median']), (264, 1240000, 365, 510000))
        self.assertEqual(ds['ytd_published']['honolulu']['current']['sf_sales'], 1950)
        self.assertEqual(ds['ytd_published']['honolulu']['current']['sf_median'], 1190000)
        self.assertEqual(ds['ytd_published']['hawaii']['current']['sf_sales'], 1335)
        self.assertEqual(ds['ytd_sum_same_months']['hawaii']['current']['sf_sales']['value'], 1317)
        gap = [g for g in ds['ytd_published_vs_sum_of_months'] if g['county'] == 'hawaii' and g['field'] == 'sf_sales'
               and g['through'] == '2026-08'][0]
        self.assertEqual(gap['difference'], 18)
        self.assertIsNone(ds['statewide_median'])
        self.assertEqual(ds['state_counts']['months']['2026-08'], {'sf_sales': 509, 'condo_sales': 479})
        self.assertIsNone(ds['published_summary']['state']['values']['sf_median']['month'])
        for row in ds['state_counts']['months'].values():
            self.assertTrue(set(row) <= {'sf_sales', 'condo_sales'})
        latest = d.read_json(os.path.join(self.state, 'last_good', d.LATEST_FILE))
        self.assertIsNone(latest['statewide_median'])
        self.assertEqual(latest['counties']['kauai']['month']['condo_median'], 999999)

    def test_output_is_deterministic(self):
        self.run_edition(AUG, '2026-10-09T02:00:00Z')
        first = self.last_good_bytes()
        self.state = os.path.join(self.tmp, 'state2')
        self.run_edition(AUG, '2026-10-09T02:00:00Z')
        self.assertEqual(first, self.last_good_bytes())


class Releases(Base):
    def test_real_july_to_august_revisions_are_detected(self):
        rec = self.run_edition(self.edition('jul'), '2026-09-03T00:00:00Z')
        self.assertEqual(rec['decision'], 'promoted')
        jul = self.last_good()
        self.assertIsNone(jul['release']['publication_date'])  # county titles say Aug 28, summary sheets Aug 27 / Sep 2
        self.assertEqual(jul['release']['publication_date_statements'],
                         {'county_workbook_titles': ['2026-08-28'], 'summary_sheet_titles': ['2026-08-27', '2026-09-02']})
        self.assertEqual(jul['release']['update_statements'], ['will be updated on September 2, 2026'])
        warn = {c['id'] for c in rec['checks'] if c['status'] == 'warn'}
        self.assertIn('summary.year_ago.hawaii.sf_sales', warn)  # summary 163 vs county workbook 165 for 2025-07
        self.assertEqual(jul['counties']['hawaii']['months']['2025-07']['sf_sales'], 165)

        rec = self.run_edition(AUG, '2026-10-09T02:00:00Z')
        self.assertEqual((rec['decision'], rec['new_months']), ('promoted', ['2026-08']))
        ds = self.last_good()
        got = {(x['county'], x['month'], x['field'], x['old'], x['new']) for x in ds['revisions']['vs_last_good']}
        self.assertEqual(got, REAL_REVISIONS_JUL_TO_AUG)
        self.assertTrue(ds['revisions']['review_required'])
        self.assertTrue(all(x['affects_report_snapshot'] for x in ds['revisions']['vs_last_good']))
        self.assertEqual(ds['revisions']['vs_verified_baseline'], [])

        rec = self.run_edition(AUG, '2026-10-09T08:00:00Z')
        self.assertEqual((rec['decision'], rec['exit_code']), ('no_change', 0))

    def test_hold_on_revision_stages_without_promoting(self):
        self.run_edition(self.edition('jul'), '2026-09-03T00:00:00Z')
        before = self.last_good_bytes()
        rec = self.run_edition(AUG, '2026-10-09T02:00:00Z', hold_on_revision=True)
        self.assertEqual(rec['decision'], 'staged_for_review')
        self.assertEqual(before, self.last_good_bytes())
        self.assertTrue(os.path.exists(os.path.join(self.state, rec['candidate_file'])))

    def test_older_release_is_rejected(self):
        self.run_edition(AUG, '2026-10-09T02:00:00Z')
        before = self.last_good_bytes()
        rec = self.run_edition(self.edition('jul'), '2026-10-10T00:00:00Z')
        self.assertEqual((rec['decision'], rec['exit_code']), ('older_than_last_good', 2))
        self.assertEqual(before, self.last_good_bytes())

    def test_same_period_republication_is_promoted_with_its_revision(self):
        self.run_edition(AUG, '2026-10-09T02:00:00Z')
        data = read(os.path.join(AUG, '2026-08-maui.xlsx'))
        r = row_of(data, '2026-07')
        folder = self.edition('aug', edit={'2026-08-maui.xlsx': lambda b: patch_xlsx(b, {COL['sf_sales'] + str(r): 56})})
        rec = self.run_edition(folder, '2026-10-12T00:00:00Z')
        self.assertEqual(rec['decision'], 'promoted_republication')
        self.assertEqual([(x['county'], x['month'], x['field'], x['old'], x['new']) for x in self.last_good()['revisions']['vs_last_good']],
                         [('maui', '2026-07', 'sf_sales', 55, 56)])


class FailuresKeepLastGood(Base):
    def setUp(self):
        super().setUp()
        self.run_edition(self.edition('jul'), '2026-09-03T00:00:00Z')
        self.before = self.last_good_bytes()

    def assertKept(self, rec, decision, code):
        self.assertEqual((rec['decision'], rec['exit_code']), (decision, code))
        self.assertTrue(rec['kept_previous_last_good'])
        self.assertEqual(self.before, self.last_good_bytes())

    def test_publisher_page_unreachable(self):
        rec = d.run(d.DirectoryFetcher(AUG, now='2026-10-09T02:00:00Z',
                                       overrides={d.PUBLISHER_PAGE: TimeoutError('timed out')}),
                    self.state, baseline=BASELINE, now='2026-10-09T02:00:00Z')
        self.assertKept(rec, 'fetch_failed', 3)

    def test_county_workbook_server_error(self):
        rec = d.run(d.DirectoryFetcher(AUG, now='2026-10-09T02:00:00Z', overrides={URL + '2026-08-kauai.xlsx': 503}),
                    self.state, baseline=BASELINE, now='2026-10-09T02:00:00Z')
        self.assertKept(rec, 'fetch_failed', 3)

    def test_county_link_missing_from_page(self):
        folder = self.edition('aug', page=lambda p: re.sub(r'<a [^>]*2026-08-kauai\.xlsx[^>]*>.*?</a>', '', p, flags=re.S))
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'discovery_failed', 3)

    def test_corrupt_workbook(self):
        folder = self.edition('aug', edit={'2026-08-maui.xlsx': lambda b: b'<html>Service unavailable</html>'})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'validation_failed', 2)

    def test_renamed_housing_column(self):
        folder = self.edition('aug', edit={'2026-08-honolulu.xlsx': lambda b: patch_xlsx(b, {'AO3': 'Single-family sales'})})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'validation_failed', 2)
        self.assertIn('layout', [c for c in rec['checks'] if c['id'] == 'parse.county.honolulu'][0]['detail'])

    def test_text_in_a_value_cell(self):
        data = read(os.path.join(AUG, '2026-08-hawaii.xlsx'))
        r = row_of(data, '2026-05')
        folder = self.edition('aug', edit={'2026-08-hawaii.xlsx': lambda b: patch_xlsx(b, {COL['sf_sales'] + str(r): 'pending'})})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'validation_failed', 2)

    def test_latest_month_incomplete(self):
        data = read(os.path.join(AUG, '2026-08-maui.xlsx'))
        r = row_of(data, '2026-08')
        folder = self.edition('aug', page=lambda p: re.sub(r'<a [^>]*mei-08p-2026\.xlsx[^>]*>.*?</a>', '', p, flags=re.S),
                              edit={'2026-08-maui.xlsx': lambda b: patch_xlsx(b, {COL['condo_median'] + str(r): '(NA)'})})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'incomplete', 2)
        rejected = d.read_json(os.path.join(self.state, rec['candidate_file']))
        self.assertIsNone(rejected['counties']['maui']['months']['2026-08']['condo_median'])

    def test_current_month_disagrees_with_dbedt_summary(self):
        data = read(os.path.join(AUG, '2026-08-honolulu.xlsx'))
        r = row_of(data, '2026-08')
        folder = self.edition('aug', edit={'2026-08-honolulu.xlsx': lambda b: patch_xlsx(b, {COL['sf_sales'] + str(r): 265})})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'validation_failed', 2)

    def test_mass_change_to_history_is_treated_as_a_layout_problem(self):
        data = read(os.path.join(AUG, '2026-08-honolulu.xlsx'))
        months = ['%04d-%02d' % (y, m) for y in range(2010, 2015) for m in range(1, 13)]
        parsed = d.parse_county_workbook(data, 'honolulu')['months']
        edits = {COL['sf_sales'] + str(row_of(data, k)): parsed[k]['sf_sales'] + 1 for k in months}
        folder = self.edition('aug', edit={'2026-08-honolulu.xlsx': lambda b: patch_xlsx(b, edits)})
        rec = self.run_edition(folder, '2026-10-09T02:00:00Z')
        self.assertKept(rec, 'validation_failed', 2)
        rec = self.run_edition(folder, '2026-10-09T03:00:00Z', accept_mass_revision=True)
        self.assertEqual(rec['decision'], 'promoted')
        self.assertEqual(len(self.last_good()['revisions']['vs_verified_baseline']), 60)


if __name__ == '__main__':
    unittest.main()
